from celery import shared_task
from celery.exceptions import SoftTimeLimitExceeded
from workers.ollama_client import verificar_ollama as _verificar_ollama_client, chat_stream as _chat_stream_client, chat_complete as _chat_complete_client
from django.conf import settings
from asgiref.sync import async_to_sync
from channels.layers import get_channel_layer
import logging
import time
import requests
import json

logger = logging.getLogger(__name__)


_OLLAMA_BASE_URL = 'https://santiagoherazo.ddns.net:11435'
_OLLAMA_MODEL    = 'llama3.1:8b'


try:
    from adapters.realtime.chatbot_logger import ids_log, IDS, new_trace_id
except ImportError:
    import enum

    class IDS(str, enum.Enum):
        WS_IN = 'WS_IN'; WS_OUT = 'WS_OUT'; API = 'API'; TASK = 'TASK'
        OLLAMA = 'OLLAMA'; CHANNEL = 'CHANNEL'; AUTH = 'AUTH'; ERROR = 'ERROR'
        BROKER = 'BROKER'; BOOT = 'BOOT'; CELERY = 'CELERY'

    def ids_log(category, conv_id=None, level='info', **kwargs):
        parts = [f'[CHATBOT-IDS-FALLBACK] [{category}]']
        if conv_id:
            parts.append(f'conv={conv_id}')
        parts += [f'{k}={v}' for k, v in kwargs.items()]
        getattr(logger, level)(' '.join(parts))

    def new_trace_id(): return 'na'


@shared_task(
    bind=True,
    queue='default',
    max_retries=2,
    time_limit=480,
    soft_time_limit=420,
    acks_late=True,
)
def procesar_mensaje_chatbot(self, conversacion_id, contenido, msg_id=None):


    from apps.chatbot.models import Conversacion, MensajeConversacion

    task_start = time.monotonic()
    celery_id  = self.request.id
    retries    = self.request.retries


    trace_id   = (self.request.headers or {}).get('trace_id', None)

    ids_log(IDS.TASK, conv_id=conversacion_id,
            trace_id=trace_id,
            msg='task_received',
            celery_id=celery_id,
            retries=retries,
            msg_id=msg_id or 'none',
            contenido_len=len(contenido))

    try:

        try:
            conv = Conversacion.objects.select_related(
                'expediente__cliente', 'expediente__tipo_auditoria',
                'usuario_interno',
            ).get(id=conversacion_id)
            ids_log(IDS.TASK, conv_id=conversacion_id,
                    msg='conversacion_loaded',
                    expediente=str(conv.expediente_id) if conv.expediente_id else 'none',
                    usuario=str(conv.usuario_interno_id) if conv.usuario_interno_id else 'none')
        except Conversacion.DoesNotExist:
            ids_log(IDS.ERROR, conv_id=conversacion_id, level='error',
                    msg='conversacion_not_found',
                    hint='UUID no existe en BD — posiblemente se borró antes del procesamiento')
            return


        _enviar_ws_evento(conversacion_id, 'chatbot_typing', '')


        qs = conv.mensajes.order_by('-fecha')
        if msg_id:
            qs = qs.exclude(pk=msg_id)
        else:
            qs = qs.exclude(contenido=contenido, rol='USUARIO')

        msgs      = list(qs.values('rol', 'contenido')[:10])
        historial = list(reversed(msgs))

        ids_log(IDS.TASK, conv_id=conversacion_id, trace_id=trace_id,
                msg='historial_built',
                mensajes_en_historial=len(historial))


        sistema           = _construir_sistema(conv)
        respuesta, tokens = _llamar_ollama_streaming(
            sistema, historial, contenido, conversacion_id
        )


        MensajeConversacion.objects.create(
            conversacion=conv,
            rol='ASISTENTE',
            contenido=respuesta,
            tokens_usados=tokens,
        )

        elapsed = time.monotonic() - task_start
        ids_log(IDS.TASK, conv_id=conversacion_id, trace_id=trace_id,
                msg='task_completed',
                tokens=tokens,
                respuesta_len=len(respuesta),
                elapsed_seconds=f'{elapsed:.1f}')

        _enviar_ws_evento(conversacion_id, 'chatbot_done', respuesta)
        verificar_escalamiento_chatbot.delay(str(conversacion_id))

    except SoftTimeLimitExceeded:
        elapsed = time.monotonic() - task_start
        ids_log(IDS.ERROR, conv_id=conversacion_id, level='warning',
                msg='soft_time_limit_exceeded',
                elapsed_seconds=f'{elapsed:.1f}',
                hint='Ollama tardó >420s — sin GPU o modelo demasiado grande')
        _enviar_ws_evento(
            conversacion_id, 'chatbot_error',
            'El modelo tardó demasiado. Intenta de nuevo — ya está cargado en memoria.',
        )

    except Exception as exc:
        elapsed = time.monotonic() - task_start
        ids_log(IDS.ERROR, conv_id=conversacion_id, level='error',
                msg='task_exception',
                exc_type=type(exc).__name__,
                exc_detail=str(exc),
                retries=retries,
                max_retries=self.max_retries,
                elapsed_seconds=f'{elapsed:.1f}')
        logger.error('Error chatbot conv=%s: %s', conversacion_id, exc, exc_info=True)
        countdown = 30 * (2 ** self.request.retries)
        if self.request.retries < self.max_retries:
            ids_log(IDS.TASK, conv_id=conversacion_id,
                    msg='task_retry_scheduled',
                    countdown_seconds=countdown,
                    attempt=retries + 1)
            raise self.retry(exc=exc, countdown=countdown)
        _enviar_ws_evento(
            conversacion_id, 'chatbot_error',
            'El asistente no pudo procesar tu mensaje. Verifica que Ollama esté activo.',
        )


def _contexto_expediente_activo(conv, get_contexto_expediente) -> str:

    if not conv.expediente_id:
        return ''
    ctx = get_contexto_expediente(str(conv.expediente_id))
    if not ctx:
        return ''
    seccion = (
        f"\n\n=== EXPEDIENTE ACTIVO ==="
        f"\nNúmero: {ctx['numero']} | Cliente: {ctx['cliente']} (NIT: {ctx['nit']})"
        f"\nTipo: {ctx['tipo']} | Estado: {ctx['estado']}"
        f"\nAvance: {ctx['avance']:.0f}% ({ctx['fases_completadas']}/{ctx['fases_total']} fases)"
        f"\nHallazgos abiertos — Críticos: {ctx['hallazgos_criticos']} | "
        f"Mayores: {ctx['hallazgos_mayores']} | Menores: {ctx['hallazgos_menores']}"
        f"\nDocumentos pendientes: {ctx['docs_pendientes']}"
    )
    if ctx.get('auditor_lider'):
        seccion += f"\nAuditor Líder: {ctx['auditor_lider']}"
    return seccion


def _contexto_expedientes_usuario(usuario_id, rol, get_expedientes_usuario) -> str:

    if not (usuario_id and rol in ('SUPERVISOR', 'ASESOR', 'AUDITOR', 'AUXILIAR', 'REVISOR')):
        return ''
    exps = get_expedientes_usuario(usuario_id, rol or '')
    if not exps:
        return ''
    seccion = f"\n\n=== TUS EXPEDIENTES ACTIVOS ({len(exps)}) ==="
    for e in exps[:5]:
        seccion += (
            f"\n• {e['numero']} — {e['cliente']} ({e['tipo']}) "
            f"| {e['estado']} | Avance: {e['avance']:.0f}%"
        )
    if len(exps) > 5:
        seccion += f"\n... y {len(exps)-5} más."
    return seccion


def _contexto_certs_por_vencer(usuario_id, rol, get_certs_por_vencer) -> str:

    uid_certs = usuario_id if rol == 'ASESOR' else None
    certs = get_certs_por_vencer(uid_certs)
    if not certs:
        return ''
    seccion = "\n\n=== CERTIFICACIONES POR VENCER (próximos 30 días) ==="
    for c in certs:
        seccion += (
            f"\n• {c['numero']} — {c['cliente']} | Vence: {c['vence']} "
            f"({c['dias_restantes']} días)"
        )
    return seccion


def _construir_sistema(conv):


    from workers.chat_context import (
        get_contexto_expediente,
        get_expedientes_usuario,
        get_certs_por_vencer,
        get_stats_globales,
    )

    rol = None
    usuario_id = None
    if conv.usuario_interno:
        rol        = getattr(conv.usuario_interno, 'rol', None)
        usuario_id = str(conv.usuario_interno.id)

    rol_contexto = {
        'SUPERVISOR': 'Eres el asistente del Supervisor. Tienes acceso total al sistema AuditCore: expedientes, clientes, hallazgos, certificaciones, usuarios y configuración.',
        'ASESOR':     'Eres el asistente del Asesor Comercial. Tu especialidad es la captación y seguimiento de clientes, coordinación del proceso de caracterización y presentación de propuestas.',
        'AUDITOR':    'Eres el asistente del Auditor. Tu especialidad es la ejecución de auditorías, normas ISO 9001 e ISO 27001, registro de hallazgos, carga de evidencias, checklists y emisión de certificaciones.',
        'AUXILIAR':   'Eres el asistente del Auxiliar de Auditoría. Tu especialidad es la elaboración de informes preliminares y comentarios de mejora sobre procedimientos en curso.',
        'REVISOR':    'Eres el asistente del Revisor. Tienes acceso de lectura a expedientes, hallazgos, documentos y certificaciones para revisión y control de calidad.',
        'CLIENTE':    'Eres el asistente del cliente de AuditCore. Puedes consultar el estado de tus auditorías, certificaciones e informes. Para detalles técnicos, contacta a tu asesor asignado.',
    }.get(rol or '', 'Eres AuditBot, asistente de AuditCore — plataforma de auditorías.')

    base = (
        f"{rol_contexto} "
        "Responde SIEMPRE en español, de forma concisa y directa. "
        "Cuando el usuario pregunte sobre expedientes, hallazgos o certificaciones, "
        "usa los datos reales del sistema que se incluyen en este contexto. "
        "Para preguntas irrelevantes al negocio de auditorías, indícalo brevemente."
    )

    base += _contexto_expediente_activo(conv, get_contexto_expediente)
    base += _contexto_expedientes_usuario(usuario_id, rol, get_expedientes_usuario)
    base += _contexto_certs_por_vencer(usuario_id, rol, get_certs_por_vencer)


    if rol == 'SUPERVISOR':
        stats = get_stats_globales()
        if stats:
            base += (
                f"\n\n=== RESUMEN GLOBAL ==="
                f"\nExpedientes activos: {stats.get('expedientes_activos', '?')}"
                f" | Hallazgos críticos abiertos: {stats.get('hallazgos_criticos', '?')}"
                f" | Clientes activos: {stats.get('clientes_activos', '?')}"
            )

    if conv.usuario_interno:
        u = conv.usuario_interno
        base += f"\n\nUsuario actual: {u.nombre_completo} ({u.get_rol_display()})"

    return base


def _verificar_ollama(base_url: str, conv_id: str | None = None) -> tuple[bool, str]:


    ids_log(IDS.OLLAMA, conv_id=conv_id,
            msg='preflight_check',
            url=f'{base_url}/api/tags')
    t0 = time.monotonic()
    try:
        r = requests.get(f'{base_url}/api/tags', timeout=5)
        r.raise_for_status()
        try:
            modelos = [m.get('name', '?') for m in r.json().get('models', [])]
        except Exception:
            modelos = ['(no parseable)']
        ids_log(IDS.OLLAMA, conv_id=conv_id,
                msg='preflight_ok',
                latency_ms=f'{(time.monotonic()-t0)*1000:.0f}',
                modelos_disponibles=','.join(modelos) or 'ninguno')
        return True, ''
    except requests.exceptions.ConnectionError:
        ids_log(IDS.OLLAMA, conv_id=conv_id, level='error',
                msg='preflight_failed',
                reason='connection_error',
                url=base_url,
                hint='docker ps | grep ollama')
        return False, (
            'El servicio Ollama no está disponible. '
            'Verifica que el contenedor esté corriendo: docker ps | grep ollama'
        )
    except requests.exceptions.Timeout:
        ids_log(IDS.OLLAMA, conv_id=conv_id, level='error',
                msg='preflight_failed',
                reason='timeout_5s',
                hint='Ollama puede estar iniciando — espera ~30s')
        return False, 'Ollama no responde (timeout 5s). Puede estar iniciando.'
    except Exception as e:
        ids_log(IDS.OLLAMA, conv_id=conv_id, level='error',
                msg='preflight_failed',
                reason=type(e).__name__,
                detail=str(e))
        return False, f'Error al contactar Ollama: {e}'


def _llamar_ollama_streaming(sistema, historial, mensaje_usuario, conversacion_id):


    base_url = getattr(settings, 'OLLAMA_BASE_URL', _OLLAMA_BASE_URL)
    model    = getattr(settings, 'OLLAMA_MODEL',    _OLLAMA_MODEL)

    ok, err_msg = _verificar_ollama_client(conv_id=conversacion_id)
    if not ok:
        ids_log(IDS.ERROR, conv_id=conversacion_id, level='error',
                msg='ollama_preflight_failed_returning_error',
                err_msg=err_msg)
        return err_msg, 0

    ids_log(IDS.OLLAMA, conv_id=conversacion_id,
            msg='stream_start',
            model=model,
            messages_count=len(historial) + 2,
            url=f'{base_url}/api/chat')

    messages = [{'role': 'system', 'content': sistema}]
    for m in historial:
        rol = 'assistant' if m['rol'] == 'ASISTENTE' else 'user'
        messages.append({'role': rol, 'content': m['contenido']})
    messages.append({'role': 'user', 'content': mensaje_usuario})

    t_stream_start  = time.monotonic()
    chunks_enviados = 0

    try:
        with requests.post(
            f'{base_url}/api/chat',
            json={
                'model':      model,
                'messages':   messages,
                'stream':     True,
                'keep_alive': -1,
                'options': {
                    'temperature':    0.7,
                    'num_predict':    768,
                    'num_ctx':        2048,
                    'num_gpu':        99,
                    'num_thread':     0,
                    'repeat_penalty': 1.1,
                },
            },
            stream=True,
            timeout=390,
        ) as response:
            response.raise_for_status()
            ids_log(IDS.OLLAMA, conv_id=conversacion_id,
                    msg='stream_http_ok',
                    status=response.status_code,
                    ttfb_ms=f'{(time.monotonic()-t_stream_start)*1000:.0f}')

            texto_completo   = []
            buffer_chunk     = []
            tokens_generados = 0

            for line in response.iter_lines():
                if not line:
                    continue
                try:
                    data = json.loads(line.decode('utf-8'))
                except (json.JSONDecodeError, UnicodeDecodeError):
                    continue

                chunk = data.get('message', {}).get('content', '')
                if chunk:
                    texto_completo.append(chunk)
                    buffer_chunk.append(chunk)

                    buffer_str = ''.join(buffer_chunk)
                    if len(buffer_str) >= 15 or any(c in buffer_str for c in '.!?;\n'):
                        _enviar_ws_evento(conversacion_id, 'chatbot_token', buffer_str)
                        chunks_enviados += 1
                        buffer_chunk = []

                if data.get('done', False):
                    if buffer_chunk:
                        _enviar_ws_evento(
                            conversacion_id, 'chatbot_token', ''.join(buffer_chunk)
                        )
                        chunks_enviados += 1
                    tokens_generados = (
                        data.get('eval_count', 0) +
                        data.get('prompt_eval_count', 0)
                    )
                    break

            elapsed_stream = time.monotonic() - t_stream_start
            respuesta = ''.join(texto_completo).strip()

            ids_log(IDS.OLLAMA, conv_id=conversacion_id,
                    msg='stream_done',
                    tokens=tokens_generados,
                    chunks_ws_enviados=chunks_enviados,
                    respuesta_len=len(respuesta),
                    stream_seconds=f'{elapsed_stream:.1f}')

            if not respuesta:
                ids_log(IDS.ERROR, conv_id=conversacion_id, level='warning',
                        msg='ollama_empty_response',
                        hint='Respuesta vacía — modelo incorrecto o prompt demasiado largo')
                return 'El asistente no pudo generar una respuesta. Intenta de nuevo.', 0

            return respuesta, tokens_generados

    except requests.exceptions.ConnectionError:
        ids_log(IDS.ERROR, conv_id=conversacion_id, level='error',
                msg='ollama_connection_error_during_stream',
                base_url=base_url,
                hint='Contenedor Ollama caído durante el streaming')
        return (
            'El asistente no está disponible. '
            'Verifica que el servicio Ollama esté corriendo.', 0
        )
    except requests.exceptions.Timeout:
        elapsed = time.monotonic() - t_stream_start
        ids_log(IDS.ERROR, conv_id=conversacion_id, level='error',
                msg='ollama_stream_timeout',
                model=model,
                elapsed_seconds=f'{elapsed:.1f}',
                hint='Con GPU <60s; sin GPU puede exceder 390s')
        return (
            'El asistente tardó demasiado. '
            'Intenta de nuevo — el modelo debería estar en memoria.', 0
        )
    except requests.exceptions.HTTPError as e:
        status_code = e.response.status_code
        ids_log(IDS.ERROR, conv_id=conversacion_id, level='error',
                msg='ollama_http_error',
                status=status_code,
                body=e.response.text[:200])
        if status_code == 404:
            return (
                f'Modelo "{model}" no encontrado en Ollama. '
                'Ejecuta: docker exec -it ollama-1 ollama pull ' + model, 0
            )
        return f'Error del asistente ({status_code}). Intenta de nuevo.', 0
    except Exception as e:
        ids_log(IDS.ERROR, conv_id=conversacion_id, level='error',
                msg='ollama_unexpected_error',
                exc_type=type(e).__name__,
                detail=str(e))
        logger.error('Error inesperado en Ollama: %s', e, exc_info=True)
        return 'Error inesperado. Intenta de nuevo.', 0


def _enviar_ws_evento(conversacion_id, event_type, contenido):


    group_name = f'chatbot_{conversacion_id}'
    t0 = time.monotonic()
    try:
        channel_layer = get_channel_layer()

        if channel_layer is None:
            ids_log(IDS.ERROR, conv_id=str(conversacion_id), level='error',
                    msg='channel_layer_none_in_worker',
                    event=event_type,
                    hint='CHANNEL_LAYERS no configurado o Redis no disponible desde el worker')
            return

        ids_log(IDS.CHANNEL, conv_id=str(conversacion_id),
                msg='group_send_attempt',
                event=event_type,
                group=group_name,
                backend=type(channel_layer).__name__,
                contenido_len=len(str(contenido)))

        async_to_sync(channel_layer.group_send)(
            group_name,
            {
                'type':      event_type,
                'contenido': contenido,
            },
        )

        ids_log(IDS.CHANNEL, conv_id=str(conversacion_id),
                msg='group_send_ok',
                event=event_type,
                group=group_name,
                latency_ms=f'{(time.monotonic()-t0)*1000:.0f}')

    except Exception as e:
        ids_log(IDS.ERROR, conv_id=str(conversacion_id), level='error',
                msg='group_send_failed',
                event=event_type,
                group=group_name,
                exc_type=type(e).__name__,
                detail=str(e),
                hint='Redis caído o cliente desconectado antes de recibir')
        logger.error('Error WS %s: %s', event_type, e)


@shared_task(queue='default')
def verificar_escalamiento_chatbot(conversacion_id):

    PALABRAS = [
        'no entiendo', 'no me ayuda', 'no sirve', 'no funciona',
        'hablar con humano', 'agente', 'persona real', 'ayuda real',
    ]
    try:
        from apps.chatbot.models import Conversacion
        conv = Conversacion.objects.select_related(
            'expediente__ejecutivo', 'usuario_interno'
        ).get(id=conversacion_id)

        ultimos = list(conv.mensajes.filter(rol='USUARIO').order_by('-fecha')[:3])
        if len(ultimos) < 3:
            return
        texto = ' '.join(m.contenido.lower() for m in ultimos)
        if not any(p in texto for p in PALABRAS):
            return
        if conv.expediente and conv.expediente.ejecutivo:
            from django.core.mail import send_mail
            ej = conv.expediente.ejecutivo
            if not ej.email:
                logger.warning('Escalamiento chatbot: ejecutivo %s sin email', ej.id)
                return
            nombre_usuario = (
                conv.usuario_interno.nombre_completo
                if conv.usuario_interno else 'Usuario desconocido'
            )
            send_mail(
                subject=f'[AuditCore] Escalamiento chatbot — {conv.expediente.numero_expediente}',
                message=(
                    f'El usuario {nombre_usuario} '
                    f'necesita asistencia humana en {conv.expediente.numero_expediente}.\n'
                    f'Revisa el historial en el sistema.'
                ),
                from_email=settings.DEFAULT_FROM_EMAIL,
                recipient_list=[ej.email],
                fail_silently=True,
            )
            ids_log(IDS.TASK, conv_id=str(conversacion_id),
                    msg='escalamiento_email_sent',
                    ejecutivo_email=ej.email)
    except Exception as e:
        ids_log(IDS.ERROR, conv_id=str(conversacion_id), level='error',
                msg='escalamiento_error',
                detail=str(e))


import os
import mimetypes

try:
    from adapters.realtime.auditlog import (
        alog, AL,
        log_doc_analysis_start, log_doc_analysis_done, log_task_start,
        log_task_done, log_task_error,
    )
except ImportError:
    def alog(*a, **kw):
        pass
    def log_doc_analysis_start(*a, **kw):
        pass
    def log_doc_analysis_done(*a, **kw):
        pass
    def log_task_start(*a, **kw):
        pass
    def log_task_done(*a, **kw):
        pass
    def log_task_error(*a, **kw):
        pass


_TIPOS_SEGUROS = {
    'application/pdf',
    'application/msword',
    'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
    'application/vnd.ms-excel',
    'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    'text/plain',
    'text/csv',
    'image/jpeg',
    'image/png',
    'image/gif',
    'image/webp',
}


_EXTENSIONES_PELIGROSAS = {
    '.exe', '.bat', '.sh', '.ps1', '.cmd', '.com', '.msi', '.dll',
    '.vbs', '.js', '.jar', '.py', '.rb', '.php', '.pl', '.scr',
    '.pif', '.reg', '.lnk', '.hta', '.wsf', '.inf', '.iso',
}


_MAX_TEXTO_ANALISIS = 6000


def _verificar_seguridad_archivo(nombre: str, mime_type: str, tamanio_bytes: int) -> tuple[bool, str]:


    ext = os.path.splitext(nombre.lower())[1]
    if ext in _EXTENSIONES_PELIGROSAS:
        return False, f'Extensión bloqueada por seguridad: {ext}. No se permiten archivos ejecutables.'

    if mime_type and mime_type not in _TIPOS_SEGUROS:
        return False, f'Tipo de archivo no permitido: {mime_type}. Solo PDF, Word, Excel, imágenes y texto plano.'

    limite_mb = 50
    if tamanio_bytes > limite_mb * 1024 * 1024:
        return False, f'Archivo demasiado grande ({tamanio_bytes // (1024*1024)} MB). Máximo permitido: {limite_mb} MB.'

    return True, ''


def _extraer_pdf(ruta_archivo: str) -> str:
    try:
        import pypdf
        reader = pypdf.PdfReader(ruta_archivo)
        paginas = []
        for i, page in enumerate(reader.pages[:20]):
            t = page.extract_text() or ''
            if t.strip():
                paginas.append(f'[Pág {i+1}] {t}')
        return '\n'.join(paginas)
    except Exception as e:
        return f'(No se pudo extraer texto del PDF: {e})'


def _extraer_word(ruta_archivo: str) -> str:
    try:
        import docx
        doc = docx.Document(ruta_archivo)
        return '\n'.join(p.text for p in doc.paragraphs if p.text.strip())
    except Exception as e:
        return f'(No se pudo extraer texto del Word: {e})'


def _extraer_excel(ruta_archivo: str) -> str:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(ruta_archivo, read_only=True, data_only=True)
        filas = []
        for sheet in wb.worksheets[:3]:
            filas.append(f'[Hoja: {sheet.title}]')
            for row in list(sheet.iter_rows(values_only=True))[:50]:
                fila_str = ' | '.join(str(c) if c is not None else '' for c in row)
                if fila_str.strip(' |'):
                    filas.append(fila_str)
        return '\n'.join(filas)
    except Exception as e:
        return f'(No se pudo extraer datos del Excel: {e})'


def _extraer_texto_plano(ruta_archivo: str) -> str:
    try:
        with open(ruta_archivo, 'r', encoding='utf-8', errors='replace') as f:
            return f.read()
    except Exception as e:
        return f'(No se pudo leer el archivo de texto: {e})'


def _extraer_csv(ruta_archivo: str) -> str:
    try:
        import csv
        filas = []
        with open(ruta_archivo, 'r', encoding='utf-8', errors='replace') as f:
            reader = csv.reader(f)
            for i, row in enumerate(reader):
                if i >= 100:
                    filas.append(f'... ({i} filas totales, mostrando primeras 100)')
                    break
                filas.append(' | '.join(row))
        return '\n'.join(filas)
    except Exception as e:
        return f'(No se pudo leer el CSV: {e})'


def _extraer_texto_documento(ruta_archivo: str, mime_type: str) -> str:


    ruta_lower = ruta_archivo.lower()
    texto = ''
    try:
        if mime_type == 'application/pdf' or ruta_lower.endswith('.pdf'):
            texto = _extraer_pdf(ruta_archivo)
        elif mime_type in (
            'application/vnd.openxmlformats-officedocument.wordprocessingml.document',
            'application/msword',
        ) or ruta_lower.endswith(('.docx', '.doc')):
            texto = _extraer_word(ruta_archivo)
        elif mime_type in (
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            'application/vnd.ms-excel',
        ) or ruta_lower.endswith(('.xlsx', '.xls')):
            texto = _extraer_excel(ruta_archivo)
        elif mime_type == 'text/plain' or ruta_lower.endswith('.txt'):
            texto = _extraer_texto_plano(ruta_archivo)
        elif mime_type == 'text/csv' or ruta_lower.endswith('.csv'):
            texto = _extraer_csv(ruta_archivo)
        elif mime_type and mime_type.startswith('image/'):
            texto = '(Archivo de imagen — análisis basado en metadatos y nombre de archivo)'
    except Exception as e:
        texto = f'(Error inesperado extrayendo contenido: {e})'

    return texto[:_MAX_TEXTO_ANALISIS] if texto else '(Sin contenido extraíble)'


def _construir_prompt_analisis(nombre_archivo: str, mime_type: str, tamanio_bytes: int,
                               texto_doc: str, pregunta_usuario: str,
                               seguro: bool, es_admin_o_lider: bool) -> str:

    extension  = os.path.splitext(nombre_archivo)[1].upper() or 'desconocido'
    tamanio_kb = tamanio_bytes // 1024
    pregunta_base = pregunta_usuario.strip() if pregunta_usuario.strip() else (
        'Analiza este documento en el contexto de una auditoría.'
    )
    nivel_detalle = (
        'Dado que el usuario es Administrador o Auditor Líder, también incluye una evaluación técnica detallada.'
        if es_admin_o_lider else
        'Usa un lenguaje claro y accesible para el usuario.'
    )
    return f"""
El usuario ha subido el siguiente documento para análisis:

📄 **Archivo:** {nombre_archivo}
📦 **Tipo:** {mime_type} ({extension}) | **Tamaño:** {tamanio_kb} KB
{'✅ **Verificación de seguridad:** Pasó' if seguro else ''}

**Pregunta/solicitud del usuario:** {pregunta_base}

---

**CONTENIDO DEL DOCUMENTO:**
{texto_doc}

---

Por favor realiza las siguientes evaluaciones del documento en el contexto de auditoría ISO:

1. **Veracidad y consistencia**: ¿El documento parece auténtico y consistente? ¿Hay inconsistencias, fechas incoherentes o datos que parezcan alterados?

2. **Completitud**: ¿Falta algún campo obligatorio, firma, fecha, número de versión, sello, o sección requerida?

3. **Conformidad**: Para auditorías ISO 9001 / ISO 27001, ¿cumple con los requisitos típicos del tipo de documento que parece ser?

4. **Riesgos identificados**: ¿Hay algo que deba revisarse antes de aprobarlo o subirlo al servidor?

5. **Recomendaciones**: Qué acciones concretas se recomiendan (aprobar, rechazar, solicitar correcciones, etc.).

{nivel_detalle}

Responde en español con formato claro usando secciones numeradas.


Llama a Ollama para analizar un documento. Retorna (respuesta, tokens)."""
    try:
        response = requests.post(
            f'{base_url}/api/chat',
            json={
                'model': model,
                'messages': [
                    {'role': 'system', 'content': contexto_rol},
                    {'role': 'user',   'content': prompt_analisis},
                ],
                'stream': False,
                'keep_alive': -1,
                'options': {
                    'temperature':    0.3,
                    'num_predict':    1200,
                    'num_ctx':        4096,
                    'num_gpu':        99,
                    'repeat_penalty': 1.1,
                },
            },
            timeout=230,
        )
        response.raise_for_status()
        data      = response.json()
        respuesta = data.get('message', {}).get('content', '').strip()
        tokens    = data.get('eval_count', 0) + data.get('prompt_eval_count', 0)
        if not respuesta:
            respuesta = 'No se pudo generar el análisis del documento. Intenta de nuevo.'
        return respuesta, tokens
    except requests.exceptions.Timeout:
        return 'El análisis tardó demasiado. Intenta con un documento más pequeño.', 0
    except Exception as e:
        ids_log(IDS.ERROR, conv_id=conversacion_id, level='error',
                msg='ollama_error_doc_analysis', detail=str(e))
        return f'Error al analizar el documento: {e}', 0


    self,
    conversacion_id: str,
    documento_id: str,
    nombre_archivo: str,
    mime_type: str,
    tamanio_bytes: int,
    ruta_archivo: str,
    pregunta_usuario: str = '',
):


    from apps.chatbot.models import Conversacion, MensajeConversacion
    from apps.documentos.models import DocumentoExpediente

    task_start = time.monotonic()
    trace_id   = (self.request.headers or {}).get('trace_id', new_trace_id())

    ids_log(IDS.TASK, conv_id=conversacion_id,
            trace_id=trace_id,
            msg='analizar_documento_task_received',
            documento_id=documento_id,
            nombre=nombre_archivo,
            mime=mime_type,
            tamanio_bytes=tamanio_bytes)

    try:

        try:
            conv = Conversacion.objects.select_related(
                'expediente__cliente', 'expediente__tipo_auditoria',
                'usuario_interno',
            ).get(id=conversacion_id)
        except Conversacion.DoesNotExist:
            ids_log(IDS.ERROR, conv_id=conversacion_id, level='error',
                    msg='conversacion_not_found_for_doc_analysis')
            return


        _enviar_ws_evento(conversacion_id, 'chatbot_typing', '')


        seguro, motivo_rechazo = _verificar_seguridad_archivo(
            nombre_archivo, mime_type, tamanio_bytes
        )
        if not seguro:
            ids_log(IDS.ERROR, conv_id=conversacion_id, level='warning',
                    msg='documento_bloqueado_por_seguridad',
                    motivo=motivo_rechazo,
                    nombre=nombre_archivo)
            msg_seguridad = (
                f'⚠️ **Archivo bloqueado por seguridad**\n\n'
                f'`{nombre_archivo}` no puede ser procesado.\n\n'
                f'**Motivo:** {motivo_rechazo}\n\n'
                f'Por favor, sube un archivo en formato permitido (PDF, Word, Excel, imagen o texto).'
            )
            MensajeConversacion.objects.create(
                conversacion=conv,
                rol='ASISTENTE',
                contenido=msg_seguridad,
                tokens_usados=0,
            )
            _enviar_ws_evento(conversacion_id, 'chatbot_done', msg_seguridad)
            return


        ids_log(IDS.TASK, conv_id=conversacion_id,
                msg='extrayendo_texto_documento',
                nombre=nombre_archivo)
        log_doc_analysis_start(nombre_archivo, conversacion_id)
        texto_doc = _extraer_texto_documento(ruta_archivo, mime_type)


        rol_usuario = getattr(conv.usuario_interno, 'rol', None) if conv.usuario_interno else None
        es_admin_o_lider = rol_usuario in ('SUPERVISOR', 'AUDITOR')
        contexto_rol = _construir_sistema(conv)

        prompt_analisis = _construir_prompt_analisis(
            nombre_archivo, mime_type, tamanio_bytes,
            texto_doc, pregunta_usuario, seguro, es_admin_o_lider,
        )


        base_url = getattr(settings, 'OLLAMA_BASE_URL', _OLLAMA_BASE_URL)
        model    = getattr(settings, 'OLLAMA_MODEL', _OLLAMA_MODEL)

        ok, err_msg = _verificar_ollama_client(conv_id=conversacion_id)
        if not ok:
            _enviar_ws_evento(conversacion_id, 'chatbot_error', err_msg)
            return

        respuesta, tokens = _llamar_ollama_doc_analysis(
            base_url, model, contexto_rol, prompt_analisis, conversacion_id
        )

        respuesta_final = f'📋 **Análisis del documento:** `{nombre_archivo}`\n\n' + respuesta


        MensajeConversacion.objects.create(
            conversacion=conv,
            rol='ASISTENTE',
            contenido=respuesta_final,
            tokens_usados=tokens,
        )

        elapsed = time.monotonic() - task_start
        ids_log(IDS.TASK, conv_id=conversacion_id,
                trace_id=trace_id,
                msg='analisis_documento_completed',
                tokens=tokens,
                elapsed_seconds=f'{elapsed:.1f}',
                seguro=seguro)
        log_doc_analysis_done(nombre_archivo, conversacion_id, tokens, elapsed)

        _enviar_ws_evento(conversacion_id, 'chatbot_done', respuesta_final)

    except SoftTimeLimitExceeded:
        _enviar_ws_evento(
            conversacion_id, 'chatbot_error',
            'El análisis tardó demasiado. Intenta con un archivo más pequeño.',
        )
    except Exception as exc:
        ids_log(IDS.ERROR, conv_id=conversacion_id, level='error',
                msg='analizar_documento_task_exception',
                exc_type=type(exc).__name__,
                detail=str(exc))
        log_task_error('analizar_documento_chatbot', str(self.request.id or ''),
                       exc, self.request.retries, self.max_retries)
        logger.error('Error analizando documento conv=%s: %s', conversacion_id, exc, exc_info=True)
        if self.request.retries < self.max_retries:
            raise self.retry(exc=exc, countdown=15)
        _enviar_ws_evento(
            conversacion_id, 'chatbot_error',
            'Error al procesar el documento. Verifica que el archivo no esté dañado.',
        )


@shared_task(
    bind=True,
    queue='default',
    max_retries=1,
    time_limit=240,
    soft_time_limit=200,
    acks_late=True,
)
def analizar_formulario_bot(self, esquema_id: str, ruta_archivo: str, mime_type: str, origen: str):


    import os
    from apps.formularios.models import EsquemaFormulario, CampoFormulario

    try:
        esquema = EsquemaFormulario.objects.get(id=esquema_id)
    except EsquemaFormulario.DoesNotExist:
        return

    try:
        texto = _extraer_texto_documento(ruta_archivo, mime_type)

        prompt = f"""Analiza el siguiente formulario y extrae todos sus campos.

CONTENIDO DEL FORMULARIO:
{texto[:5000]}

Responde ÚNICAMENTE con un JSON válido con esta estructura exacta:
{{
  "descripcion": "descripción breve del formulario en 1-2 oraciones",
  "campos": [
    {{
      "nombre": "nombre_campo_sin_espacios",
      "etiqueta": "Etiqueta visible para el usuario",
      "tipo": "TEXTO|NUMERO|FECHA|LISTA|BOOLEANO|ARCHIVO",
      "obligatorio": true|false,
      "orden": 1,
      "opciones": [],
      "ayuda": "texto de ayuda opcional"
    }}
  ]
}}

Tipos válidos: TEXTO (respuesta abierta), NUMERO (valor numérico), FECHA (fecha/hora),
LISTA (selección de opciones — incluye opciones en el array), BOOLEANO (sí/no), ARCHIVO (adjunto).
Para campos de tipo LISTA, incluye las opciones disponibles en el array "opciones".
No incluyas campos de firma, logo o membrete — solo campos de datos.

